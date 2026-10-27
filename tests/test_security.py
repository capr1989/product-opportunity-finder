import ast
import io
from pathlib import Path
import socket

import httpx
from openpyxl import load_workbook
import pytest

from pof.demo import demo_products
from pof.exporting import csv_bytes, excel_bytes
from pof.models import Product
from pof.security import MAX_RESPONSE, SafeHTTP, safe_path, spreadsheet_safe, validate_url


@pytest.fixture
def public_dns(monkeypatch):
    monkeypatch.setattr(socket, "getaddrinfo", lambda *a, **kw: [(2, 1, 6, "", ("93.184.216.34", 443))])


@pytest.mark.parametrize("url", ["file:///etc/passwd", "http://example.com", "https://user:secret@example.com",
    "https://example.com:444", "https://evil.com", "https://example.com/#fragment"])
def test_bad_url(public_dns, url):
    with pytest.raises(ValueError):
        validate_url(url, "example.com")


@pytest.mark.parametrize("ip", ["127.0.0.1", "10.0.0.1", "169.254.169.254", "::1", "192.168.1.1", "0.0.0.0"])
def test_private_dns(monkeypatch, ip):
    monkeypatch.setattr(socket, "getaddrinfo", lambda *a, **kw: [(2, 1, 6, "", (ip, 443))])
    with pytest.raises(ValueError):
        validate_url("https://example.com/", "example.com")


def test_redirect_and_mime(public_dns):
    for response in [httpx.Response(302, headers={"location": "http://127.0.0.1"}),
                     httpx.Response(200, headers={"content-type": "application/octet-stream"})]:
        client = SafeHTTP(0.5, httpx.MockTransport(lambda r: response))
        with pytest.raises(ValueError):
            client.fetch("https://example.com", "example.com")


def test_dns_pinning_and_success(public_dns):
    def handler(request):
        assert request.url.host == "93.184.216.34"
        assert request.headers["host"] == "example.com"
        assert request.extensions["sni_hostname"] == "example.com"
        return httpx.Response(200, headers={"content-type": "application/json"}, content=b"[]")
    assert SafeHTTP(0.5, httpx.MockTransport(handler)).fetch("https://example.com/feed", "example.com") == b"[]"


def test_oversized_response(public_dns):
    transport = httpx.MockTransport(lambda r: httpx.Response(200, headers={"content-type": "application/json"},
                                                             content=b"x" * (MAX_RESPONSE + 1)))
    with pytest.raises(ValueError):
        SafeHTTP(0.5, transport).fetch("https://example.com", "example.com")


@pytest.mark.parametrize("name", ["../escape.csv", "C:\\escape.csv", "x/escape.csv", "evil.exe", "x:stream.csv"])
def test_path_traversal(tmp_path, name):
    with pytest.raises(ValueError):
        safe_path(tmp_path, name)


@pytest.mark.parametrize("value", ["=HYPERLINK(\"https://evil.com\")", " +SUM(1,1)", "\t@evil", "-cmd"])
def test_formula_injection(value):
    assert spreadsheet_safe(value).startswith("'")
    assert "'" in csv_bytes([{"value": value}]).decode("utf-8-sig")
    p = Product("test", value)
    workbook = load_workbook(io.BytesIO(excel_bytes([p], [], {})))
    assert all(cell.data_type != "f" for sheet in workbook for row in sheet for cell in row)


def test_export_structure():
    workbook = load_workbook(io.BytesIO(excel_bytes(demo_products(), [], {})))
    assert len(workbook.sheetnames) == 8
    assert workbook["All products"].max_row == 13


def test_no_shell_or_unsafe_deserialization():
    for path in Path("pof").rglob("*.py"):
        tree = ast.parse(path.read_text())
        for node in ast.walk(tree):
            if isinstance(node, ast.Call):
                assert not (isinstance(node.func, ast.Name) and node.func.id in {"eval", "exec"})
                assert not any(k.arg == "shell" and isinstance(k.value, ast.Constant) and k.value.value for k in node.keywords)
    config = Path(".streamlit/config.toml").read_text()
    assert 'address = "127.0.0.1"' in config
    assert 'server.address=127.0.0.1' in Path("Launch.cmd").read_text()


def test_browser_children_do_not_inherit_app_credentials(monkeypatch):
    from pof.browser import browser_env
    monkeypatch.setenv("GOOGLE_ADS_REFRESH_TOKEN", "private-value")
    monkeypatch.setenv("BESTBUY_API_KEY", "private-value")
    monkeypatch.setenv("HTTPS_PROXY", "private-proxy")
    result = browser_env()
    assert "GOOGLE_ADS_REFRESH_TOKEN" not in result
    assert "BESTBUY_API_KEY" not in result
    assert "HTTPS_PROXY" not in result
