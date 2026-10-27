"""Memory-only CSV/XLSX exports with formula-injection protection."""
import csv
import io
import json

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from pof.security import spreadsheet_safe

SHEETS = ["All products", "Top opportunities", "Supplier comparisons", "Competitor comparisons",
          "Profit calculations", "Keyword research", "Source references", "Assumptions and missing data"]


def flatten(row: dict) -> dict:
    return {str(k): spreadsheet_safe(json.dumps(v, ensure_ascii=False) if isinstance(v, (dict, list)) else v)
            for k, v in row.items()}


def csv_bytes(rows: list[dict]) -> bytes:
    stream = io.StringIO(newline="")
    clean = [flatten(row) for row in rows]
    if clean:
        writer = csv.DictWriter(stream, fieldnames=list(dict.fromkeys(k for r in clean for k in r)))
        writer.writeheader()
        writer.writerows(clean)
    return stream.getvalue().encode("utf-8-sig")


def excel_bytes(products, opportunities, assumptions, keywords=None) -> bytes:
    records = [p.record() for p in products]
    data = [records, opportunities[:100], [r for r in opportunities if r["role"] == "supplier"],
            [r for r in opportunities if r["role"] == "retailer"],
            [{k: v for k, v in r.items() if k not in {"source_record", "retail_record", "breakdown"}}
             for r in opportunities], keywords or [],
            [{"source": p.source, "url": p.url, "collected_at": p.collected_at, "demo": p.demo,
              "field_states": {k: p.state(k) for k in p.record()}} for p in products],
            [{"type": "assumptions", **assumptions}] + [{"type": "missing/stale", "title": p.title,
              "fields": {k: p.state(k) for k in p.record() if p.state(k) != "verified"}} for p in products]]
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, rows in zip(SHEETS, data):
        sheet = workbook.create_sheet(name)
        clean = [flatten(r) for r in rows]
        if not clean:
            sheet.append(["No data available"])
            continue
        keys = list(dict.fromkeys(k for r in clean for k in r))
        sheet.append(keys)
        for row in clean:
            sheet.append([row.get(key) for key in keys])
        for cell in sheet[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = PatternFill("solid", fgColor="164E63")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = 24
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
